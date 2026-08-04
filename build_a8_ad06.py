#!/usr/bin/env python3
"""ATMOS_A8_IDEF0_AD06_vNext.pptx - native editable A8 IDEF0 (A81/A82/A83)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

BLACK=RGBColor(0,0,0); WHITE=RGBColor(0xFF,0xFF,0xFF); FONT='Arial'
prs=Presentation(); prs.slide_width=Inches(11); prs.slide_height=Inches(8.5)
sl=prs.slides.add_slide(prs.slide_layouts[6]); SH=sl.shapes
def I(v): return Inches(v)

BL,BT,BR,BB = 0.30,0.38,10.70,8.16
BOX={'A81':(1.70,3.55,3.70,5.15),'A82':(4.55,3.55,6.60,5.15),'A83':(7.50,3.55,9.70,5.15)}
NAME={'A81':'Compare Predicted and Observed Weather',
      'A82':'Develop and Validate Candidate Model Improvements',
      'A83':'Package and Archive Assessment and Learning Artifacts'}
CX={'A81':2.70,'A82':5.575,'A83':8.60}          # control / mechanism stem x (box centres)
TOPY,BOTY = 3.55,5.15
CORR = 2.95                                     # O4 corridor lane
conns=[]
def port(x,y):
    s=SH.add_shape(MSO_SHAPE.OVAL,I(x)-Emu(9144),I(y)-Emu(9144),Emu(18288),Emu(18288))
    s.fill.background(); s.line.fill.background(); s.shadow.inherit=False; s._pt=(x,y); return s
def seg(p,q,arrow):
    (x1,y1),(x2,y2)=p._pt,q._pt
    c=SH.add_connector(MSO_CONNECTOR.STRAIGHT,I(x1),I(y1),I(x2),I(y2))
    c.begin_connect(p,0); c.end_connect(q,0)
    x=c._element.spPr.find(qn('a:xfrm')); o=x.find(qn('a:off')); e=x.find(qn('a:ext'))
    o.set('x',str(int(I(min(x1,x2))))); o.set('y',str(int(I(min(y1,y2)))))
    e.set('cx',str(int(I(abs(x2-x1))))); e.set('cy',str(int(I(abs(y2-y1)))))
    if x2<x1: x.set('flipH','1')
    if y2<y1: x.set('flipV','1')
    c.line.color.rgb=BLACK; c.line.width=Pt(1.25); c.shadow.inherit=False
    if arrow:
        ln=c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn('a:tailEnd'),{'type':'triangle','w':'med','len':'med'}))
    conns.append(c)
def route(pts,arrow=True):
    ps=[port(*p) for p in pts]
    for i in range(len(ps)-1): seg(ps[i],ps[i+1],arrow and i==len(ps)-2)

# ---- INPUTS: F1/F2/F4 -> three separate left ports on A81 ----
route([(BL,3.85),(1.70,3.85)])
route([(BL,4.43),(1.70,4.43)])
route([(BL,5.01),(1.70,5.01)])
# ---- CONTROLS: one per box, straight down into the top face ----
for k in ('A81','A82','A83'): route([(CX[k],1.94),(CX[k],TOPY)])
# ---- MECHANISMS: one per box, straight up into the bottom face ----
for k in ('A81','A82','A83'): route([(CX[k],6.86),(CX[k],BOTY)])
# ---- A81 output trunk (IER-23) + three branches ----
route([(3.70,4.20),(3.95,4.20)],arrow=False)                 # trunk off A81 right face
route([(3.95,4.20),(4.55,4.20)])                             # branch 1 -> A82 left
route([(3.95,4.20),(3.95,CORR)],arrow=False)                 # riser to corridor
route([(3.95,CORR),(BR,CORR)])                               # branch 3 -> right boundary (O4)
route([(7.05,CORR),(7.05,3.95),(7.50,3.95)])                 # branch 2 -> A83 left
# ---- A82 internal output -> A83 ----
route([(6.60,4.60),(7.50,4.60)])                             # A8-F1
# ---- A83 output ----
route([(9.70,4.30),(BR,4.30)])                               # O5

# ---- BOXES ----
for k,(l,t,r,b) in BOX.items():
    s=SH.add_shape(MSO_SHAPE.RECTANGLE,I(l),I(t),I(r-l),I(b-t))
    s.fill.solid(); s.fill.fore_color.rgb=WHITE
    s.line.color.rgb=BLACK; s.line.width=Pt(1.25); s.shadow.inherit=False
    tf=s.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    tf.margin_left=Pt(6); tf.margin_right=Pt(6); tf.margin_top=Pt(3); tf.margin_bottom=Pt(15)
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    run=p.add_run(); run.text=NAME[k]
    run.font.size=Pt(11); run.font.bold=True; run.font.name=FONT; run.font.color.rgb=BLACK
    n=SH.add_textbox(I(r-0.60),I(b-0.26),I(0.54),I(0.20))
    ntf=n.text_frame; ntf.word_wrap=False
    for m in ('margin_left','margin_right','margin_top','margin_bottom'): setattr(ntf,m,0)
    np_=ntf.paragraphs[0]; np_.alignment=PP_ALIGN.RIGHT
    nr=np_.add_run(); nr.text=k
    nr.font.size=Pt(12); nr.font.bold=True; nr.font.name=FONT; nr.font.color.rgb=BLACK

# ---- FRAME + LABELS ----
fr=SH.add_shape(MSO_SHAPE.RECTANGLE,I(BL),I(BT),I(BR-BL),I(BB-BT))
fr.fill.background(); fr.line.color.rgb=BLACK; fr.line.width=Pt(1.0); fr.shadow.inherit=False
def lab(t,x,y,w,h,sz=10,b=False,a=PP_ALIGN.LEFT):
    tb=SH.add_textbox(I(x),I(y),I(w),I(h)); tf=tb.text_frame; tf.word_wrap=True
    tf.margin_left=Pt(1); tf.margin_right=Pt(1); tf.margin_top=Pt(0); tf.margin_bottom=Pt(0)
    p=tf.paragraphs[0]; p.alignment=a
    r=p.add_run(); r.text=t
    r.font.size=Pt(sz); r.font.bold=b; r.font.name=FONT; r.font.color.rgb=BLACK
lab('A8 — Decompose Assess Weather-Model Performance and Learn',BL,0.50,BR-BL,0.36,17,True,PP_ALIGN.CENTER)
lab('A8',9.98,7.72,0.60,0.24,12,True,PP_ALIGN.RIGHT)
# inputs (each label sits in the clear band directly above its own arrow)
lab('F1 Qualified Observations (IER-03)',0.36,3.26,1.26,0.56)
lab('F2 Local Model Estimate (IER-05)',0.36,3.85,1.26,0.56)
lab('F4 Fused Weather State (IER-12)',0.36,4.43,1.26,0.56)
# controls
lab('C81 Assessment Criteria',1.70,1.56,2.00,0.36,10,False,PP_ALIGN.CENTER)
lab('C82 Model Governance',4.55,1.56,2.05,0.36,10,False,PP_ALIGN.CENTER)
lab('C83 Archival Constraints',7.50,1.56,2.20,0.36,10,False,PP_ALIGN.CENTER)
# mechanisms
lab('M81 Assessment Analytics',1.70,6.92,2.00,0.40,10,False,PP_ALIGN.CENTER)
lab('M82 Model Development Resources',4.55,6.92,2.05,0.40,10,False,PP_ALIGN.CENTER)
lab('M83 Archive Export Services',7.50,6.92,2.20,0.40,10,False,PP_ALIGN.CENTER)
# flows / outputs (each product labelled once)
lab('O4 Weather Assessment (IER-23)',9.78,2.36,0.84,0.57)
lab('O5 Learning Archive (IER-24)',9.78,3.70,0.84,0.57)
lab('A8-F1 Candidate Improvement Package',6.64,4.70,0.86,0.76)
prs.save('/home/user/Default/ATMOS_A8_IDEF0_AD06_vNext.pptx')
print('saved shapes=',len(sl.shapes._spTree),'connectors=',len(conns))
