#!/usr/bin/env python3
"""ATMOS_A8_IDEF0_AD06_vNext_R1.pptx — built from a blank slide.
Visible text is drawn ONLY from the exclusive inventory below."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

# ---------- exclusive visible-text inventory ----------
TITLE = 'A8 — Decompose Assess Weather-Model Performance and Learn'
NODE  = 'A8'
FN    = {'A81':'Compare Predicted and Observed Weather',
         'A82':'Develop and Validate Candidate Model Improvements',
         'A83':'Package and Archive Assessment and Learning Artifacts'}
IN1,IN2,IN3 = ('F1 Qualified Observations (IER-03)',
               'F2 Local Model Estimate (IER-05)',
               'F4 Fused Weather State (IER-12)')
C1,C2,C3    = 'C81 Assessment Criteria','C82 Model Governance','C83 Archival Constraints'
FLOW        = 'A8-F1 Candidate Improvement Package'
OUT4,OUT5   = 'O4 Weather Assessment (IER-23)','O5 Learning Archive (IER-24)'
M1,M2,M3    = 'M81 Assessment Analytics','M82 Model Development Resources','M83 Archive Export Services'
INVENTORY = {TITLE,NODE,FLOW,OUT4,OUT5,IN1,IN2,IN3,C1,C2,C3,M1,M2,M3,
             *FN.keys(), *FN.values()}

BLACK=RGBColor(0,0,0); WHITE=RGBColor(0xFF,0xFF,0xFF); FONT='Arial'
prs=Presentation(); prs.slide_width=Inches(11); prs.slide_height=Inches(8.5)
sl=prs.slides.add_slide(prs.slide_layouts[6])          # blank layout
SH=sl.shapes
def I(v): return Inches(v)

# ---------- geometry ----------
BL,BT,BR,BB = 0.30,0.38,10.70,8.16
BOX={'A81':(1.70,3.55,3.70,5.15),'A82':(4.55,3.55,6.60,5.15),'A83':(7.50,3.55,9.70,5.15)}
STEM={'A81':2.70,'A82':5.575,'A83':8.60}               # control/mechanism stem x = box centre
TOPY,BOTY = 3.55,5.15
CORR = 2.95                                            # O4 branch corridor
IN_Y = (3.85,4.43,5.01)                                # three A81 left ports
TRUNK_Y = 4.20                                         # A81 right face -> A82 left face
A83_O4_Y, A83_F1_Y = 3.95, 4.60                        # two A83 left ports
O5_Y = 4.30
conns=[]

def port(x,y):
    s=SH.add_shape(MSO_SHAPE.OVAL,I(x)-Emu(9144),I(y)-Emu(9144),Emu(18288),Emu(18288))
    s.fill.background(); s.line.fill.background(); s.shadow.inherit=False; s._pt=(x,y); return s

def seg(p,q,arrow):
    (x1,y1),(x2,y2)=p._pt,q._pt
    c=SH.add_connector(MSO_CONNECTOR.STRAIGHT,I(x1),I(y1),I(x2),I(y2))
    c.begin_connect(p,0); c.end_connect(q,0)           # native endpoint attachment
    x=c._element.spPr.find(qn('a:xfrm')); o=x.find(qn('a:off')); e=x.find(qn('a:ext'))
    o.set('x',str(int(I(min(x1,x2))))); o.set('y',str(int(I(min(y1,y2)))))
    e.set('cx',str(int(I(abs(x2-x1))))); e.set('cy',str(int(I(abs(y2-y1)))))
    if x2<x1: x.set('flipH','1')
    if y2<y1: x.set('flipV','1')
    c.line.color.rgb=BLACK; c.line.width=Pt(1.25); c.shadow.inherit=False
    if arrow:
        ln=c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn('a:tailEnd'),{'type':'triangle','w':'med','len':'med'}))
    conns.append(c)

def route(pts,arrow=True):
    ps=[port(*p) for p in pts]
    for i in range(len(ps)-1): seg(ps[i],ps[i+1],arrow and i==len(ps)-2)

# ---------- connectors first (so boxes and labels sit on top) ----------
for y in IN_Y: route([(BL,y),(1.70,y)])                             # F1,F2,F4 -> A81 left
for k in BOX: route([(STEM[k],1.94),(STEM[k],TOPY)])                # C81,C82,C83 -> tops
for k in BOX: route([(STEM[k],6.86),(STEM[k],BOTY)])                # M81,M82,M83 -> bottoms
route([(3.70,TRUNK_Y),(3.95,TRUNK_Y)],arrow=False)                  # O4 trunk off A81
route([(3.95,TRUNK_Y),(4.55,TRUNK_Y)])                              #   branch -> A82 left
route([(3.95,TRUNK_Y),(3.95,CORR)],arrow=False)                     #   riser to corridor
route([(3.95,CORR),(BR,CORR)])                                      #   branch -> right boundary
route([(7.05,CORR),(7.05,A83_O4_Y),(7.50,A83_O4_Y)])                #   branch -> A83 left
route([(6.60,A83_F1_Y),(7.50,A83_F1_Y)])                            # A8-F1 A82 -> A83 left
route([(9.70,O5_Y),(BR,O5_Y)])                                      # O5 -> right boundary

# ---------- function boxes ----------
for k,(l,t,r,b) in BOX.items():
    s=SH.add_shape(MSO_SHAPE.RECTANGLE,I(l),I(t),I(r-l),I(b-t))
    s.fill.solid(); s.fill.fore_color.rgb=WHITE
    s.line.color.rgb=BLACK; s.line.width=Pt(1.25); s.shadow.inherit=False
    tf=s.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    tf.margin_left=Pt(6); tf.margin_right=Pt(6); tf.margin_top=Pt(3); tf.margin_bottom=Pt(15)
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    r_=p.add_run(); r_.text=FN[k]
    r_.font.size=Pt(11); r_.font.bold=True; r_.font.name=FONT; r_.font.color.rgb=BLACK
    n=SH.add_textbox(I(r-0.60),I(b-0.26),I(0.54),I(0.20))
    ntf=n.text_frame; ntf.word_wrap=False
    for m in ('margin_left','margin_right','margin_top','margin_bottom'): setattr(ntf,m,0)
    np_=ntf.paragraphs[0]; np_.alignment=PP_ALIGN.RIGHT
    nr=np_.add_run(); nr.text=k
    nr.font.size=Pt(12); nr.font.bold=True; nr.font.name=FONT; nr.font.color.rgb=BLACK

# ---------- border + labels ----------
fr=SH.add_shape(MSO_SHAPE.RECTANGLE,I(BL),I(BT),I(BR-BL),I(BB-BT))
fr.fill.background(); fr.line.color.rgb=BLACK; fr.line.width=Pt(1.0); fr.shadow.inherit=False
def lab(t,x,y,w,h,sz=10,b=False,a=PP_ALIGN.LEFT):
    assert t in INVENTORY, 'text not in inventory: %r'%t
    tb=SH.add_textbox(I(x),I(y),I(w),I(h)); tf=tb.text_frame; tf.word_wrap=True
    tf.margin_left=Pt(1); tf.margin_right=Pt(1); tf.margin_top=Pt(0); tf.margin_bottom=Pt(0)
    p=tf.paragraphs[0]; p.alignment=a
    r=p.add_run(); r.text=t
    r.font.size=Pt(sz); r.font.bold=b; r.font.name=FONT; r.font.color.rgb=BLACK
lab(TITLE,BL,0.50,BR-BL,0.36,17,True,PP_ALIGN.CENTER)
lab(NODE,9.98,7.72,0.60,0.24,12,True,PP_ALIGN.RIGHT)
lab(IN1,0.36,3.26,1.26,0.56); lab(IN2,0.36,3.85,1.26,0.56); lab(IN3,0.36,4.43,1.26,0.56)
lab(C1,1.70,1.56,2.00,0.36,10,False,PP_ALIGN.CENTER)
lab(C2,4.55,1.56,2.05,0.36,10,False,PP_ALIGN.CENTER)
lab(C3,7.50,1.56,2.20,0.36,10,False,PP_ALIGN.CENTER)
lab(M1,1.70,6.92,2.00,0.40,10,False,PP_ALIGN.CENTER)
lab(M2,4.55,6.92,2.05,0.40,10,False,PP_ALIGN.CENTER)
lab(M3,7.50,6.92,2.20,0.40,10,False,PP_ALIGN.CENTER)
lab(OUT4,9.78,2.36,0.84,0.57)                       # O4 labelled once, at the boundary
lab(OUT5,9.78,3.70,0.84,0.57)
lab(FLOW,6.64,4.70,0.86,0.76)
prs.save('/home/user/Default/ATMOS_A8_IDEF0_AD06_vNext_R1.pptx')
print('saved; shapes=%d connectors=%d'%(len(sl.shapes._spTree),len(conns)))
