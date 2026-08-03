#!/usr/bin/env python3
"""Render a native pptx slide to PNG from its real shape geometry (for visual QA)."""
import sys, textwrap
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from pptx import Presentation
from pptx.oxml.ns import qn
EMU=914400.0
fn=sys.argv[1]; out=sys.argv[2]
prs=Presentation(fn); W=prs.slide_width/EMU; H=prs.slide_height/EMU
sl=prs.slides[0]
fig,ax=plt.subplots(figsize=(W,H),dpi=200)
ax.set_xlim(0,W); ax.set_ylim(0,H); ax.invert_yaxis(); ax.axis('off')
ax.add_patch(Rectangle((0,0),W,H,fc='white',ec='none',zorder=0))
def geom(el):
    x=el.find(qn('p:spPr')).find(qn('a:xfrm'))
    o=x.find(qn('a:off')); e=x.find(qn('a:ext'))
    X=int(o.get('x'))/EMU; Y=int(o.get('y'))/EMU
    CX=int(e.get('cx'))/EMU; CY=int(e.get('cy'))/EMU
    return X,Y,CX,CY,x.get('flipH')=='1',x.get('flipV')=='1'
for sp in sl.shapes:
    el=sp._element; tag=el.tag.split('}')[-1]
    if tag=='cxnSp':
        X,Y,CX,CY,fH,fV=geom(el)
        bx,ex=(X+CX,X) if fH else (X,X+CX); by,ey=(Y+CY,Y) if fV else (Y,Y+CY)
        ln=el.find(qn('p:spPr')).find(qn('a:ln'))
        arr=ln is not None and ln.find(qn('a:tailEnd')) is not None
        ax.plot([bx,ex],[by,ey],color='k',lw=0.85,solid_capstyle='butt',zorder=2)
        if arr:
            d=((ex-bx)**2+(ey-by)**2)**0.5 or 1
            ux,uy=(ex-bx)/d,(ey-by)/d
            ax.annotate('',xy=(ex,ey),xytext=(ex-ux*0.09,ey-uy*0.09),
                arrowprops=dict(arrowstyle='-|>',color='k',lw=0.8,
                                mutation_scale=7),zorder=3)
        continue
    if tag!='sp': continue
    spPr=el.find(qn('p:spPr'))
    if spPr is None or spPr.find(qn('a:xfrm')) is None: continue
    X,Y,CX,CY,_,_=geom(el)
    txt=sp.text_frame.text if sp.has_text_frame else ''
    solid=spPr.find(qn('a:solidFill')) is not None
    lnE=spPr.find(qn('a:ln'))
    has_line=lnE is not None and lnE.find(qn('a:noFill')) is None
    if solid: ax.add_patch(Rectangle((X,Y),CX,CY,fc='white',ec='none',zorder=4))
    if has_line: ax.add_patch(Rectangle((X,Y),CX,CY,fc='none',ec='k',lw=1.0,zorder=5))
    if txt.strip():
        szs=[int(r.get('sz')) for r in el.iter(qn('a:rPr')) if r.get('sz')]
        pt=(szs[0]/100) if szs else 10
        bold=any(r.get('b')=='1' for r in el.iter(qn('a:rPr')))
        algn=[p.get('algn') for p in el.iter(qn('a:pPr'))]
        al=algn[0] if algn else 'l'
        fs=pt*0.86
        cpl=max(6,int(CX*72/(fs*0.55)))
        lines=[]
        for L in txt.split('\n'): lines+= textwrap.wrap(L,cpl) or ['']
        disp='\n'.join(lines)
        if al=='ctr': hx,ha=X+CX/2,'center'
        elif al=='r': hx,ha=X+CX,'right'
        else: hx,ha=X+0.02,'left'
        va,vy=('center',Y+CY/2) if solid else ('top',Y)
        ax.text(hx,vy,disp,fontsize=fs,ha=ha,va=va,fontweight='bold' if bold else 'normal',
                family='DejaVu Sans',color='k',zorder=6,linespacing=1.05)
plt.subplots_adjust(0,0,1,1); plt.savefig(out,dpi=200,facecolor='white'); print('wrote',out)
