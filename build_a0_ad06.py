#!/usr/bin/env python3
"""ATMOS_A0_IDEF0_AD06_vNext.pptx - native editable A0 IDEF0 (A1-A6 + A8, no A7)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

BLACK=RGBColor(0,0,0); WHITE=RGBColor(0xFF,0xFF,0xFF); FONT='Arial'
prs=Presentation(); prs.slide_width=Inches(11); prs.slide_height=Inches(8.5)
sl=prs.slides.add_slide(prs.slide_layouts[6]); SH=sl.shapes
def I(v): return Inches(v)

BL,BT,BR,BB = 0.30,0.38,10.70,8.16
BOX={'A1':(1.44,3.52,2.60,4.40),'A2':(2.90,3.52,4.05,4.40),'A3':(4.42,3.52,5.56,4.40),
     'A4':(5.84,3.52,7.34,4.40),'A5':(7.60,1.96,8.90,2.82),'A6':(7.60,4.72,8.90,5.58),
     'A8':(4.40,6.40,6.30,7.26)}
NAME={'A1':'Acquire Micro-Weather Observations','A2':'Generate Local Weather State',
      'A3':'Quantify Uncertainty','A4':'Federate and Maintain Federated Weather Context',
      'A5':'Detect Threshold Crossings (Descriptive Support)',
      'A6':'Produce Mission-Tailored Weather Context',
      'A8':'Assess Weather-Model Performance and Learn'}
C4Y,C1Y,C2Y,C3Y = 1.20,1.32,1.44,1.56
M1Y,M2Y,M3Y     = 5.68,5.80,5.92
I3Y=1.72
conns=[]
def port(x,y):
    s=SH.add_shape(MSO_SHAPE.OVAL,I(x)-Emu(9144),I(y)-Emu(9144),Emu(18288),Emu(18288))
    s.fill.background(); s.line.fill.background(); s.shadow.inherit=False; s._pt=(x,y); return s
def seg(p,q,arrow):
    (x1,y1),(x2,y2)=p._pt,q._pt
    c=SH.add_connector(MSO_CONNECTOR.STRAIGHT,I(x1),I(y1),I(x2),I(y2))
    c.begin_connect(p,0); c.end_connect(q,0)
    x=c._element.spPr.find(qn('a:xfrm')); o=x.find(qn('a:off')); e=x.find(qn('a:ext'))
    o.set('x',str(int(I(min(x1,x2))))); o.set('y',str(int(I(min(y1,y2)))))
    e.set('cx',str(int(I(abs(x2-x1))))); e.set('cy',str(int(I(abs(y2-y1)))))
    if x2<x1: x.set('flipH','1')
    if y2<y1: x.set('flipV','1')
    c.line.color.rgb=BLACK; c.line.width=Pt(1.25); c.shadow.inherit=False
    if arrow:
        ln=c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn('a:tailEnd'),{'type':'triangle','w':'med','len':'med'}))
    conns.append(c)
def route(pts,arrow=True):
    ps=[port(*p) for p in pts]
    for i in range(len(ps)-1): seg(ps[i],ps[i+1],arrow and i==len(ps)-2)

# ---------------- INPUTS ----------------
route([(BL,3.72),(1.44,3.72)])                                    # I1 -> A1
route([(BL,4.10),(1.44,4.10)])                                    # I2 -> A1
route([(BL,I3Y),(7.38,I3Y),(7.38,2.16),(7.60,2.16)])              # I3 -> A5 LEFT
route([(BL,4.78),(2.74,4.78),(2.74,4.14),(2.90,4.14)])            # I4 -> A2 LEFT
# ---------------- CONTROLS ----------------
for x,y in ((1.90,C1Y),(3.86,C2Y),(6.12,C3Y),(8.34,C4Y)): route([(x,1.16),(x,y)],arrow=False)
route([(8.34,C4Y),(3.10,C4Y)],arrow=False)                        # C4 bus
route([(1.90,C1Y),(8.98,C1Y)],arrow=False)                        # C1 bus
route([(3.86,C2Y),(9.06,C2Y)],arrow=False)                        # C2 bus
route([(5.68,C3Y),(7.05,C3Y)],arrow=False)                        # C3 bus
route([(3.10,C4Y),(3.10,3.52)])                                   # C4 -> A2
route([(6.10,C2Y),(6.10,3.52)]); route([(7.05,C3Y),(7.05,3.52)])  # C2,C3 -> A4
route([(8.00,C1Y),(8.00,1.96)]); route([(8.55,C2Y),(8.55,1.96)])  # C1,C2 -> A5
route([(8.98,C1Y),(8.98,4.48),(8.00,4.48),(8.00,4.72)])           # C1 -> A6
route([(9.06,C2Y),(9.06,4.60),(8.55,4.60),(8.55,4.72)])           # C2 -> A6
route([(2.68,C1Y),(2.68,6.14),(4.62,6.14),(4.62,6.40)])           # C1 -> A8
route([(4.34,C2Y),(4.34,6.26),(5.12,6.26),(5.12,6.40)])           # C2 -> A8
route([(5.68,C3Y),(5.68,6.40)])                                   # C3 -> A8
# ---------------- INTERNAL FLOWS ----------------
route([(2.60,3.80),(2.90,3.80)])                                  # F1 A1->A2
route([(4.05,3.92),(4.42,3.92)])                                  # F2 A2->A3
route([(5.56,3.92),(5.84,3.92)])                                  # F3 A3->A4
route([(7.34,3.64),(7.42,3.64),(7.42,2.52),(7.60,2.52)])          # F4 A4->A5
route([(7.34,4.20),(7.52,4.20),(7.52,5.02),(7.60,5.02)])          # F5 A4->A6
route([(2.80,3.80),(2.80,7.00),(4.40,7.00)])                      # F1 branch -> A8
route([(4.14,3.92),(4.14,6.56),(4.40,6.56)])                      # F2 branch -> A8
route([(7.42,3.64),(7.42,6.02),(4.24,6.02),(4.24,6.78),(4.40,6.78)])  # F4 branch -> A8
# ---------------- OUTPUTS ----------------
route([(7.34,3.92),(BR,3.92)])                                    # O1 <- A4
route([(8.90,2.36),(BR,2.36)])                                    # O3 <- A5
route([(8.90,5.02),(BR,5.02)])                                    # O2 <- A6
route([(6.30,6.66),(BR,6.66)])                                    # O4 <- A8
route([(6.30,7.12),(BR,7.12)])                                    # O5 <- A8
# ---------------- MECHANISMS ----------------
route([(1.80,7.58),(1.80,M1Y)],arrow=False)
route([(3.52,7.58),(3.52,M2Y)],arrow=False)
route([(6.90,7.58),(6.90,M3Y)],arrow=False)
route([(1.72,M1Y),(9.22,M1Y)],arrow=False)
route([(3.52,M2Y),(5.25,M2Y)],arrow=False)
route([(2.30,M3Y),(9.14,M3Y)],arrow=False)
for x in (1.72,3.20,4.70,6.05): route([(x,M1Y),(x,4.40)])         # M1 -> A1..A4
route([(3.72,M2Y),(3.72,4.40)]); route([(5.25,M2Y),(5.25,4.40)])  # M2 -> A2,A3
route([(2.30,M3Y),(2.30,4.40)]); route([(7.15,M3Y),(7.15,4.40)])  # M3 -> A1,A4
route([(8.00,M1Y),(8.00,5.58)]); route([(8.55,M3Y),(8.55,5.58)])  # M1,M3 -> A6
route([(9.22,M1Y),(9.22,3.20),(8.65,3.20),(8.65,2.82)])           # M1 -> A5
route([(9.14,M3Y),(9.14,3.02),(8.82,3.02),(8.82,2.82)])           # M3 -> A5
route([(3.70,M1Y),(3.70,7.54),(5.85,7.54),(5.85,7.26)])           # M1 -> A8
route([(3.82,M2Y),(3.82,7.46),(5.35,7.46),(5.35,7.26)])           # M2 -> A8
route([(3.94,M3Y),(3.94,7.38),(4.85,7.38),(4.85,7.26)])           # M3 -> A8

# ---------------- BOXES ----------------
for k,(l,t,r,b) in BOX.items():
    s=SH.add_shape(MSO_SHAPE.RECTANGLE,I(l),I(t),I(r-l),I(b-t))
    s.fill.solid(); s.fill.fore_color.rgb=WHITE
    s.line.color.rgb=BLACK; s.line.width=Pt(1.25); s.shadow.inherit=False
    tf=s.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.TOP
    tf.margin_left=Pt(3); tf.margin_right=Pt(3); tf.margin_top=Pt(4); tf.margin_bottom=Pt(2)
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    run=p.add_run(); run.text=NAME[k]
    run.font.size=Pt(11); run.font.bold=True; run.font.name=FONT; run.font.color.rgb=BLACK
    n=SH.add_textbox(I(r-0.52),I(b-0.24),I(0.46),I(0.20))
    ntf=n.text_frame; ntf.word_wrap=False
    for m in ('margin_left','margin_right','margin_top','margin_bottom'): setattr(ntf,m,0)
    np_=ntf.paragraphs[0]; np_.alignment=PP_ALIGN.RIGHT
    nr=np_.add_run(); nr.text=k
    nr.font.size=Pt(11); nr.font.bold=True; nr.font.name=FONT; nr.font.color.rgb=BLACK

# ---------------- FRAME + LABELS ----------------
fr=SH.add_shape(MSO_SHAPE.RECTANGLE,I(BL),I(BT),I(BR-BL),I(BB-BT))
fr.fill.background(); fr.line.color.rgb=BLACK; fr.line.width=Pt(1.0); fr.shadow.inherit=False
def lab(t,x,y,w,h,sz=10,b=False,a=PP_ALIGN.LEFT):
    tb=SH.add_textbox(I(x),I(y),I(w),I(h)); tf=tb.text_frame; tf.word_wrap=True
    tf.margin_left=Pt(1); tf.margin_right=Pt(1); tf.margin_top=Pt(0); tf.margin_bottom=Pt(0)
    p=tf.paragraphs[0]; p.alignment=a
    r=p.add_run(); r.text=t
    r.font.size=Pt(sz); r.font.bold=b; r.font.name=FONT; r.font.color.rgb=BLACK
lab('A0 — Decompose Conduct Air-Focused Weather Exploitation Operations',BL,0.42,BR-BL,0.34,17,True,PP_ALIGN.CENTER)
lab('A0',9.95,7.86,0.62,0.24,12,True,PP_ALIGN.RIGHT)
lab('I1 Collocated Atmospheric Observations',0.34,3.06,1.06,0.52)
lab('I2 Peer Platform Weather Observations',0.34,4.14,1.06,0.52)
lab('I3 Mission Weather Threshold Definitions',0.34,1.26,1.30,0.42)
lab('I4 Reachback Gap-Fill and Boundary-Condition Data (IER-25)',0.34,4.84,1.30,0.66)
lab('C1 Mission objectives and operational constraints',0.38,0.80,1.85,0.36)
lab('C2 OPSEC / classification guidance',2.34,0.80,1.70,0.36)
lab('C3 Network availability and DDS QoS policies',4.30,0.80,1.85,0.36)
lab('C4 Approved Model / Parameter Baseline (IER-26)',6.40,0.80,2.00,0.36)
lab('F1 Time/Geo-Tagged, Quality-Checked Observations (IER-03)',1.48,2.90,1.12,0.56)
lab('F2 Local Micro-Weather State Estimate (IER-05)',3.16,2.90,0.88,0.56)
lab('F3 Confidence Bounds & Risk Envelopes (IER-09)',4.46,2.90,1.12,0.56)
lab('F4 Locally Fused COWP State (IER-12)',6.26,2.90,0.66,0.56)
lab('F5 Locally Fused COWP State (IER-12)',6.26,4.48,0.66,0.56)
lab('O3 Weather State Change Notifications (IER-17)',9.32,1.92,1.32,0.40)
lab('O1 Locally Fused COWP State (IER-12)',9.32,3.48,1.32,0.40)
lab('O2 Mission-Tailored COWP Excerpts (IER-18)',9.32,4.58,1.32,0.40)
lab('O4 Predicted vs. Observed Weather Assessment (IER-23)',9.32,6.18,1.32,0.44)
lab('O5 Archived Lessons Learned & Model Updates (IER-24)',9.32,7.20,1.32,0.44)
lab('M1 Platforms hosting ATMOS node compute',0.90,7.64,1.86,0.34)
lab('M2 ABLE-LBM / reduced models',3.10,7.64,1.50,0.34)
lab('M3 DDS middleware and communications links',5.60,7.64,2.26,0.34)
prs.save('/home/user/Default/ATMOS_A0_IDEF0_AD06_vNext.pptx')
print('saved shapes=',len(sl.shapes._spTree),'connectors=',len(conns))
