#!/usr/bin/env python3
"""Automated IDEF0 QA gate for the generated A0 pptx."""
import sys
from pptx import Presentation
from pptx.oxml.ns import qn
EMU=914400.0
prs=Presentation(sys.argv[1]); sl=prs.slides[0]
BOXN=('A1','A2','A3','A4','A5','A6','A8')
segs=[]; boxes={}; labels=[]; fonts=[]; frame=None
def geom(el):
    x=el.find(qn('p:spPr')).find(qn('a:xfrm')); o=x.find(qn('a:off')); e=x.find(qn('a:ext'))
    return (int(o.get('x'))/EMU,int(o.get('y'))/EMU,int(e.get('cx'))/EMU,int(e.get('cy'))/EMU,
            x.get('flipH')=='1',x.get('flipV')=='1')
for sp in sl.shapes:
    el=sp._element; tag=el.tag.split('}')[-1]
    spPr=el.find(qn('p:spPr'))
    if spPr is None or spPr.find(qn('a:xfrm')) is None: continue
    X,Y,CX,CY,fH,fV=geom(el)
    for r in el.iter(qn('a:rPr')):
        if r.get('sz'): fonts.append((int(r.get('sz'))/100, (sp.text_frame.text[:30] if sp.has_text_frame else '')))
    if tag=='cxnSp':
        bx,ex=(X+CX,X) if fH else (X,X+CX); by,ey=(Y+CY,Y) if fV else (Y,Y+CY)
        ln=spPr.find(qn('a:ln'))
        segs.append((bx,by,ex,ey, ln is not None and ln.find(qn('a:tailEnd')) is not None,
                     ln is not None and ln.find(qn('a:headEnd')) is not None)); continue
    txt=sp.text_frame.text.strip() if sp.has_text_frame else ''
    solid=spPr.find(qn('a:solidFill')) is not None
    lnE=spPr.find(qn('a:ln')); lined=lnE is not None and lnE.find(qn('a:noFill')) is None
    if solid and lined and txt in BOXN[:0]+tuple(): pass
    if solid and lined: boxes[txt]=(X,Y,X+CX,Y+CY)
    elif lined and not solid and CX>9: frame=(X,Y,X+CX,Y+CY)
    elif txt: labels.append((txt,X,Y,X+CX,Y+CY))
# map boxes by node id textbox
node_of={}
for t,x0,y0,x1,y1 in labels:
    if t in BOXN:
        for k,(a,b,c,d) in boxes.items():
            if a-0.02<=x0 and x1<=c+0.02 and b<=y0 and y1<=d+0.02: node_of[k]=t
BX={node_of.get(k,k):v for k,v in boxes.items()}
BX={k:v for k,v in BX.items() if k in BOXN}
fail=[]; warn=[]
print('slide %.2f x %.2f in'%(prs.slide_width/EMU,prs.slide_height/EMU))
print('boxes found:',{k:tuple(round(z,2) for z in v) for k,v in sorted(BX.items())})
if len(BX)!=7: fail.append('expected 7 function boxes, got %d'%len(BX))
if 'A7' in BX: fail.append('A7 present')
# 1 orthogonality + no head-end arrows
for s in segs:
    if abs(s[0]-s[2])>0.005 and abs(s[1]-s[3])>0.005: fail.append('diagonal seg %s'%(s,))
    if s[5]: fail.append('headEnd arrow on %s'%(s,))
# 2 no segment through a box interior
def through(seg,r,pad=0.012,eps=0.02):
    # true penetration only: touching / abutting edges do not count
    bx,by,ex,ey=seg[:4]; l,t,rr,b=r
    l+=pad; t+=pad; rr-=pad; b-=pad
    if abs(by-ey)<0.005:                        # horizontal
        return (t+eps)<by<(b-eps) and min(bx,ex)<(rr-eps) and max(bx,ex)>(l+eps)
    if abs(bx-ex)<0.005:                        # vertical
        return (l+eps)<bx<(rr-eps) and min(by,ey)<(b-eps) and max(by,ey)>(t+eps)
    return False
for s in segs:
    for k,r in BX.items():
        if through(s,r): fail.append('connector crosses BOX %s: %s'%(k,tuple(round(z,2) for z in s[:4])))
# 3 no segment through a text label
for s in segs:
    for t,x0,y0,x1,y1 in labels:
        if t in BOXN: continue
        if through(s,(x0,y0,x1,y1),pad=0.0):
            fail.append('connector crosses LABEL %r: %s'%(t[:34],tuple(round(z,2) for z in s[:4])))
# 4 arrowheads must land on a box face (or the drawing border for boundary ICOMs)
def onface(x,y,tol=0.03):
    for k,(l,t,r,b) in BX.items():
        if abs(x-l)<tol and t-tol<=y<=b+tol: return (k,'L')
        if abs(x-r)<tol and t-tol<=y<=b+tol: return (k,'R')
        if abs(y-t)<tol and l-tol<=x<=r+tol: return (k,'T')
        if abs(y-b)<tol and l-tol<=x<=r+tol: return (k,'B')
def inside(x,y,pad=0.03):
    for k,(l,t,r,b) in BX.items():
        if l+pad<x<r-pad and t+pad<y<b-pad: return k
from collections import Counter
faces=Counter(); nface=0; nbnd=0
for s in segs:
    if not s[4]: continue
    ex,ey=s[2],s[3]
    if inside(ex,ey): fail.append('ARROWHEAD INSIDE box %s at (%.2f,%.2f)'%(inside(ex,ey),ex,ey))
    f=onface(ex,ey)
    if f:
        nface+=1; faces[f]+=1
        horiz=abs(s[1]-s[3])<0.005; vert=abs(s[0]-s[2])<0.005
        if f[1] in 'LR' and not horiz: fail.append('non-perpendicular into %s %s'%f)
        if f[1] in 'TB' and not vert:  fail.append('non-perpendicular into %s %s'%f)
    elif frame and (abs(ex-frame[2])<0.03 or abs(ex-frame[0])<0.03): nbnd+=1
    else: fail.append('arrowhead terminates on whitespace at (%.2f,%.2f)'%(ex,ey))
# 5 fonts
small=[f for f in fonts if f[0]<10]
if small: fail.append('font < 10pt: %s'%small[:5])
# 6 everything inside the border
if frame:
    for t,x0,y0,x1,y1 in labels:
        if x0<frame[0]-0.005 or x1>frame[2]+0.005 or y0<frame[1]-0.005 or y1>frame[3]+0.005:
            warn.append('label outside border: %r'%t[:30])
print('\nICOM arrivals per box face:')
per={}
for (k,f),n in faces.items(): per.setdefault(k,{})[f]=n
for k in sorted(per): print('  ',k,per[k])
print('\narrowheads: on box faces=%d  at border=%d  total=%d'%(nface,nbnd,nface+nbnd))
print('segments=%d  labels=%d  min font=%.1fpt'%(len(segs),len(labels),min(f[0] for f in fonts)))
print('\n--- WARN (%d) ---'%len(warn))
for w in warn[:20]: print('  ',w)
print('--- FAIL (%d) ---'%len(fail))
for f in fail[:40]: print('  ',f)
print('\nRESULT:','PASS' if not fail else 'FAIL')
